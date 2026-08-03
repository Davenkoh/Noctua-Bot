"""Import the dorm's resident whitelist from an Excel sheet.

    python -m bot.roster_import residents.xlsx

The sheet needs one row of headers containing a name column, a room column
and a telegram column (anything matching ``tele``/``tag``/``handle``/
``username``); data rows follow underneath. The import replaces the whole
roster in one transaction.

Deliberately free of any ``telegram`` import — this runs from a shell while
the bot is polling, and SQLite's WAL mode lets both processes share the file.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from . import config, db, rooms, util

HANDLE_RE = re.compile(r"tele|tag|handle|username", re.IGNORECASE)
ROOM_RE = re.compile(r"room", re.IGNORECASE)
NAME_RE = re.compile(r"name", re.IGNORECASE)

NAME_LIMIT = 60  # sheets carry full legal names; residents type shorter ones
USAGE = "usage: python -m bot.roster_import <path.xlsx>"


def _cell_text(value: object) -> str:
    """Excel hands back ints, floats and dates — normalise to plain text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _column(cells: list[str], pattern: re.Pattern[str], taken: set[int]) -> int | None:
    for index, text in enumerate(cells):
        if index not in taken and pattern.search(text):
            return index
    return None


def find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    """First row carrying all three columns, plus their indices.

    The telegram column is claimed first so a "Telegram name" header does not
    get mistaken for the resident's name.
    """
    for index, cells in enumerate(rows):
        taken: set[int] = set()
        columns: dict[str, int] = {}
        for key, pattern in (("handle", HANDLE_RE), ("room", ROOM_RE), ("name", NAME_RE)):
            found = _column(cells, pattern, taken)
            if found is None:
                break
            columns[key] = found
            taken.add(found)
        else:
            return index, columns
    return None


def _get(cells: list[str], index: int) -> str:
    return cells[index] if index < len(cells) else ""


def _clean_handle(raw: str) -> str:
    return "".join(raw.split()).lstrip("@").lower()


def import_roster(path: Path) -> int:
    """Read ``path`` and replace the roster. Returns a process exit code."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        print(f"❌ Could not open {path}: {exc}")
        return 1

    try:
        sheet = workbook.worksheets[0]
        rows = [[_cell_text(value) for value in raw] for raw in sheet.iter_rows(values_only=True)]
        title = sheet.title
    finally:
        workbook.close()

    header = find_header(rows)
    if header is None:
        print("❌ No header row found. I need one row with a name column, a room")
        print("   column and a telegram column, e.g.:  Name | Room | Telegram handle")
        return 1

    header_index, columns = header
    header_cells = rows[header_index]
    print(f"Sheet {title!r} · header on row {header_index + 1}: ", end="")
    print(
        f"name={_get(header_cells, columns['name'])!r} "
        f"room={_get(header_cells, columns['room'])!r} "
        f"handle={_get(header_cells, columns['handle'])!r}"
    )

    entries: dict[str, tuple[str, str, str]] = {}
    notes: list[str] = []
    skipped = duplicates = 0

    for number, cells in enumerate(rows[header_index + 1 :], start=header_index + 2):
        raw_handle = _get(cells, columns["handle"])
        raw_name = _get(cells, columns["name"])
        raw_room = _get(cells, columns["room"])
        if not (raw_handle or raw_name or raw_room):
            continue  # blank spacer row

        handle = _clean_handle(raw_handle)
        if not handle:
            skipped += 1
            notes.append(f"row {number}: no telegram handle ({raw_name or 'unnamed'})")
            continue

        name = util.clean_name(raw_name, limit=NAME_LIMIT)
        if name is None:
            skipped += 1
            reason = (
                "name is empty"
                if not raw_name
                else f"name is longer than {NAME_LIMIT} characters"
            )
            notes.append(f"row {number}: {reason} (@{handle})")
            continue

        result = rooms.validate_room(raw_room)
        if result.needs_letter:
            skipped += 1
            notes.append(
                f"row {number}: room #{result.base} needs its unit letter A–F (@{handle})"
            )
            continue
        if not result.ok:
            skipped += 1
            notes.append(f"row {number}: {result.error} (@{handle})")
            continue

        if handle in entries:
            duplicates += 1
            notes.append(f"row {number}: duplicate handle @{handle}, this row wins")
        entries[handle] = (handle, name, result.room)

    for note in notes:
        print(f"  · {note}")

    if not entries:
        print("❌ No valid rows, so the roster was left untouched.")
        return 1

    db.init_db()
    db.roster_replace(list(entries.values()))
    print(f"✅ Imported {len(entries)} resident(s) into {config.DB_PATH}")
    print(f"   {skipped} row(s) skipped · {duplicates} duplicate(s) overwritten")
    return 0


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if args and args[0] in {"-h", "--help"}:
        print(USAGE)
        return 0
    if len(args) != 1:
        print(USAGE)
        return 2

    path = Path(args[0]).expanduser()
    if not path.is_file():
        print(f"❌ No such file: {path}")
        return 1
    return import_roster(path)


if __name__ == "__main__":
    raise SystemExit(main())
